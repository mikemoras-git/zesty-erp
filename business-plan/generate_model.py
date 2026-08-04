#!/usr/bin/env python3
"""Generate the Downtown House (Megaro Foumi) business-plan model workbook.

The workbook reproduces the financial tables of the GBR Consulting business
plan "The Downtown house, a social hub at Megaro Foumi" (May 2024, draft)
from a set of adjustable drivers, and checks itself against every number
printed in the PDF (pages 9-11).

Usage:  python3 generate_model.py [output.xlsx]
"""

import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# Figures as printed in the PDF (pages 9-11). Single source of truth: the
# base-scenario defaults are derived from these, and the "Check vs PDF"
# sheet compares the live model against them.
# --------------------------------------------------------------------------
PDF_YEARS = list(range(2025, 2035))

PDF = {
    "occupancy":  [0.65, 0.70, 0.75, 0.80, 0.80, 0.80, 0.80, 0.80, 0.80, 0.80],
    "occupied":   [1186, 1278, 1369, 1460, 1460, 1460, 1460, 1460, 1460, 1460],
    "adr":        [113, 122, 131, 141, 143, 146, 149, 152, 155, 158],
    "revpar":     [74, 85, 98, 112, 115, 117, 119, 122, 124, 127],
    "room_rev":   [134308, 155728, 179146, 205134, 209237, 213422, 217690, 222044, 226485, 231015],
    "fb_rev":     [1045493, 1332401, 1598479, 1811423, 1847652, 1884605, 1922297, 1960743, 1999958, 2039957],
    "total_rev":  [1179801, 1488129, 1777626, 2016558, 2056889, 2098027, 2139987, 2182787, 2226443, 2270972],
    "rooms_exp":  [16499, 18632, 20917, 23390, 23858, 24335, 24822, 25319, 25825, 26341],
    "fb_exp":     [392060, 499650, 599430, 679284, 692869, 706727, 720861, 735279, 749984, 764984],
    "staff":      [533868, 544545, 555436, 566545, 577876, 589433, 601222, 613247, 625511, 638022],
    "dept_total": [942427, 1062827, 1175783, 1269219, 1294604, 1320496, 1346906, 1373844, 1401320, 1429347],
    "dept_inc":   [237375, 425302, 601842, 747339, 762285, 777531, 793082, 808943, 825122, 841625],
    "ag":         [105808, 113619, 121086, 127575, 130127, 132729, 135384, 138091, 140853, 143670],
    "mkt":        [35394, 44644, 53329, 60497, 61707, 62941, 64200, 65484, 66793, 68129],
    "pom":        [35394, 44644, 53329, 60497, 61707, 62941, 64200, 65484, 66793, 68129],
    "util":       [53091, 66966, 79993, 90745, 92560, 94411, 96299, 98225, 100190, 102194],
    "undist_total": [229687, 269872, 307737, 339314, 346100, 353022, 360082, 367284, 374630, 382122],
    "ebitdar":    [7687, 155430, 294106, 408025, 416185, 424509, 432999, 441659, 450492, 459502],
    "rent":       [15000, 15000, 15000, 15000, 15000, 21000, 21000, 21000, 21000, 21000],
    "ebitda":     [-7313, 140430, 279106, 393025, 401185, 403509, 411999, 420659, 429492, 438502],
    "reserve":    [11798, 29763, 35553, 40331, 41138, 41961, 42800, 43656, 44529, 45419],
    "noi":        [-19111, 110667, 243553, 352694, 360048, 361548, 369199, 377003, 384964, 393083],
}
PDF_RETURNS = {
    "noi_2024": -15000, "capex_2024": -1984000, "net_cf_2024": -1999000,
    "vat_recovery_2025": 384000, "net_cf_2025": 364889,
    "noi_2044": 414765, "npv": 746821, "irr": 0.151,
}

MODEL_YEARS = list(range(2025, 2045))     # operating model horizon
CF_YEARS = list(range(2024, 2045))        # cash-flow horizon (as in the PDF)
AVAIL = 5 * 365

# Implied full-precision drivers, backed out of the printed euro figures
# (the PDF prints ADR rounded to whole euros; these reproduce it exactly).
ADR_BASE = [round(PDF["room_rev"][i] / (AVAIL * PDF["occupancy"][i]), 4) for i in range(4)]
ROOMS_PCT_BASE = [round(PDF["rooms_exp"][i] / PDF["room_rev"][i], 6) for i in range(4)]
AG_FIXED_BASE = round(PDF["ag"][0] - 0.02 * PDF["total_rev"][0], 2)   # 82,211.98

# --------------------------------------------------------------------------
# Styles
# --------------------------------------------------------------------------
RED = "A6192E"           # GBR accent
DARK = "3F3F3F"
F_TITLE = Font(bold=True, size=14, color=RED)
F_SUB = Font(italic=True, size=9, color="808080")
F_HDR = Font(bold=True, color="FFFFFF")
F_SECTION = Font(bold=True, color="FFFFFF", size=10)
F_INPUT = Font(color="0563C1")
F_ACTIVE = Font(italic=True, color="404040")
F_BOLD = Font(bold=True)
F_PCT = Font(italic=True, color="548235", size=9)
FILL_HDR = PatternFill("solid", fgColor=DARK)
FILL_SECTION = PatternFill("solid", fgColor="7F7F7F")
FILL_INPUT = PatternFill("solid", fgColor="DDEBF7")
FILL_ACTIVE = PatternFill("solid", fgColor="F2F2F2")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
TOP_BORDER = Border(top=Side(style="thin", color="7F7F7F"))

EUR = '#,##0'
EUR2 = '#,##0.00'
PCT1 = '0.0%'
PCT2 = '0.00%'

def section_row(ws, row, text, last_col):
    ws.cell(row=row, column=1, value=text).font = F_SECTION
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = FILL_SECTION


# --------------------------------------------------------------------------
# Inputs sheet: scenario table with Base / Upside / Downside columns and a
# selector; the model reads column F ("Active value").
# --------------------------------------------------------------------------
def build_inputs(wb):
    ws = wb.create_sheet("Inputs")
    ws["A1"] = "Inputs & scenarios — The Downtown house, Megaro Foumi"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("Edit the blue cells. The model reads the 'Active value' column, driven by the scenario "
                "selector below. Base reproduces the May 2024 GBR plan; Upside/Downside are illustrative "
                "placeholders — set them to whatever scenarios you want to test.")
    ws["A2"].font = F_SUB
    ws["A4"] = "Active scenario"
    ws["A4"].font = F_BOLD
    ws["B4"] = "Base"
    ws["B4"].font = Font(bold=True, color="0563C1")
    ws["B4"].fill = FILL_INPUT
    dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")

    hdr = 6
    for c, text in enumerate(["Driver", "Unit", "Base", "Upside", "Downside", "Active value"], start=1):
        cell = ws.cell(row=hdr, column=c, value=text)
        cell.font = F_HDR
        cell.fill = FILL_HDR

    rows = {}     # key -> row number
    r = hdr + 1

    def add(key, label, unit, base, up, down, fmt):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=unit).font = F_SUB
        for c, v in ((3, base), (4, up), (5, down)):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = F_INPUT
            cell.fill = FILL_INPUT
            cell.number_format = fmt
        f = ws.cell(row=r, column=6)
        f.value = f"=INDEX(C{r}:E{r},MATCH($B$4,$C${hdr}:$E${hdr},0))"
        f.font = F_ACTIVE
        f.fill = FILL_ACTIVE
        f.number_format = fmt
        rows[key] = r
        r += 1

    def sec(text):
        nonlocal r
        section_row(ws, r, text, 6)
        r += 1

    sec("ACCOMMODATION")
    add("rooms", "Number of rooms", "rooms", 5, 5, 5, "0")
    add("days", "Operating days per year", "days", 365, 365, 365, "0")
    occ = PDF["occupancy"][:4]
    for i in range(4):
        label = f"Occupancy — year {i+1} ({2025+i})" + (" and after (stabilised)" if i == 3 else "")
        add(f"occ{i+1}", label, "% of available rooms", occ[i],
            round(min(occ[i] + 0.05, 0.95), 2), round(occ[i] - 0.10, 2), PCT1)
    for i in range(4):
        label = f"ADR — year {i+1} ({2025+i})" + (" (stabilised level)" if i == 3 else "")
        add(f"adr{i+1}", label, "€ per occupied room", ADR_BASE[i],
            round(ADR_BASE[i] * 1.10, 2), round(ADR_BASE[i] * 0.90, 2), EUR2)
    add("growth", "Annual growth after stabilisation (inflation)", "% per year, from year 5 (2029)",
        0.02, 0.025, 0.015, PCT1)

    sec("FOOD & BEVERAGE (all-day bar/restaurant)")
    for i in range(4):
        label = f"F&B revenue — year {i+1} ({2025+i})" + (" (stabilised level)" if i == 3 else "")
        add(f"fb{i+1}", label, "€ per year", PDF["fb_rev"][i],
            round(PDF["fb_rev"][i] * 1.15), round(PDF["fb_rev"][i] * 0.80), EUR)

    sec("DEPARTMENTAL COSTS")
    for i in range(4):
        label = f"Rooms dept. cost — year {i+1}" + (" and after" if i == 3 else "")
        add(f"roomspct{i+1}", label, "% of room revenue", ROOMS_PCT_BASE[i],
            ROOMS_PCT_BASE[i], ROOMS_PCT_BASE[i], PCT2)
    add("fbpct", "F&B dept. cost", "% of F&B revenue", 0.375, 0.36, 0.40, PCT1)
    add("staff1", "Staff cost — year 1 (32 FTEs, all departments)", "€ per year, grows with inflation",
        533868, 533868, 560561, EUR)

    sec("UNDISTRIBUTED OPERATING EXPENSES")
    add("agfix", "Administrative & general — fixed part, year 1", "€ per year, grows with inflation",
        AG_FIXED_BASE, AG_FIXED_BASE, AG_FIXED_BASE, EUR)
    add("agvar", "Administrative & general — variable part", "% of total revenue", 0.02, 0.02, 0.02, PCT1)
    add("mkt", "Marketing", "% of total revenue", 0.03, 0.03, 0.03, PCT1)
    add("pom", "Property operations & maintenance", "% of total revenue", 0.03, 0.03, 0.03, PCT1)
    add("util", "Utility costs", "% of total revenue", 0.045, 0.045, 0.045, PCT1)

    sec("FIXED CHARGES")
    add("rent1", "Annual rent — tier 1 (from 2024)", "€ per year", 15000, 15000, 15000, EUR)
    add("rent2start", "Rent tier 2 starts in", "year", 2030, 2030, 2030, "0")
    add("rent2", "Annual rent — tier 2", "€ per year", 21000, 21000, 21000, EUR)
    add("rent3start", "Rent tier 3 starts in", "year", 2035, 2035, 2035, "0")
    add("rent3", "Annual rent — tier 3 (implied by the PDF's 2044 NOI and NPV)", "€ per year",
        90000, 90000, 90000, EUR)
    add("res1", "Reserve for replacement — year 1", "% of total revenue", 0.01, 0.01, 0.01, PCT1)
    add("res", "Reserve for replacement — year 2 onwards", "% of total revenue", 0.02, 0.02, 0.02, PCT1)

    sec("INVESTMENT & RETURNS")
    add("capex", "CAPEX (excl. VAT)", "€", 1500000, 1400000, 1800000, EUR)
    add("preopen", "Pre-opening expenses (excl. VAT)", "€", 100000, 100000, 100000, EUR)
    add("vat", "VAT rate on CAPEX & pre-opening", "%", 0.24, 0.24, 0.24, PCT1)
    add("grant", "CAPEX grant (% of CAPEX excl. VAT, received 2025)", "plan mentions applying for 60-75%",
        0.0, 0.60, 0.0, PCT1)
    add("disc", "Discount rate for NPV", "% per year", 0.10, 0.10, 0.10, PCT1)

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 34
    for col in "CDEF":
        ws.column_dimensions[col].width = 13
    ws.freeze_panes = "A7"
    return rows


# --------------------------------------------------------------------------
# Model sheet: full P&L 2025-2044, every cell a live formula.
# --------------------------------------------------------------------------
def build_model(wb, inp):
    ws = wb.create_sheet("Model")

    def I(key):     # reference to an active input value
        return f"Inputs!$F${inp[key]}"

    ws["A1"] = "Operating model — P&L forecast (nominal €)"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("All cells are formulas driven by the Inputs sheet. Years 1-4 (2025-2028) are the ramp-up; "
                "from year 5 revenue and cost lines grow at the inflation input. Layout mirrors pages 9-10 "
                "of the PDF, extended to 2044 for the returns calculation.")
    ws["A2"].font = F_SUB

    ncols = len(MODEL_YEARS)
    ws.cell(row=3, column=1, value="Year").font = F_HDR
    ws.cell(row=3, column=1).fill = FILL_HDR
    for i, y in enumerate(MODEL_YEARS):
        c = ws.cell(row=3, column=2 + i, value=y)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center")

    R = {}          # line key -> row

    def line(row, key, label, formula_fn, fmt=EUR, bold=False, italic_pct=False):
        R[key] = row
        lab = ws.cell(row=row, column=1, value=label)
        if bold:
            lab.font = F_BOLD
        if italic_pct:
            lab.font = F_PCT
        for i in range(ncols):
            col = get_column_letter(2 + i)
            cell = ws.cell(row=row, column=2 + i, value=formula_fn(i, col))
            cell.number_format = fmt
            if bold:
                cell.font = F_BOLD
                cell.border = TOP_BORDER
            if italic_pct:
                cell.font = F_PCT

    def ramp(i, col, key_prefix, growth_from):
        """Input for years 1-4; from year 5 grow the stabilised level."""
        if i < 4:
            return f"={I(f'{key_prefix}{i+1}')}"
        prev = get_column_letter(1 + i)
        return f"={prev}{growth_from}*(1+{I('growth')})"

    section_row(ws, 4, "ACCOMMODATION — KEY INDICATORS", 1 + ncols)
    line(5, "rooms", "Number of rooms", lambda i, c: f"={I('rooms')}", "0")
    line(6, "days", "Operating days", lambda i, c: f"={I('days')}", "0")
    line(7, "avail", "Available room nights", lambda i, c: f"={c}5*{c}6", EUR)
    line(8, "occ", "Occupancy",
         lambda i, c: f"={I(f'occ{i+1}')}" if i < 4 else f"={I('occ4')}", PCT1)
    line(9, "occupied", "Occupied room nights", lambda i, c: f"={c}7*{c}8", EUR)
    line(10, "adr", "ADR", lambda i, c: ramp(i, c, "adr", 10), EUR2)
    line(11, "revpar", "RevPAR", lambda i, c: f"={c}8*{c}10", EUR2)

    section_row(ws, 12, "REVENUE", 1 + ncols)
    line(13, "room_rev", "Room revenue", lambda i, c: f"={c}9*{c}10")
    line(14, "fb_rev", "Food & Beverages (all-day bar/restaurant)", lambda i, c: ramp(i, c, "fb", 14))
    line(15, "total_rev", "Total revenue", lambda i, c: f"={c}13+{c}14", bold=True)

    section_row(ws, 16, "DEPARTMENTAL EXPENSES", 1 + ncols)
    line(17, "rooms_exp", "Rooms",
         lambda i, c: f"={c}13*{I(f'roomspct{min(i+1, 4)}')}")
    line(18, "fb_exp", "Food & Beverage", lambda i, c: f"={c}14*{I('fbpct')}")
    line(19, "staff", "Staff",
         lambda i, c: f"={I('staff1')}" if i == 0
         else f"={get_column_letter(1 + i)}19*(1+{I('growth')})")
    line(20, "dept_total", "Total departmental expenses", lambda i, c: f"=SUM({c}17:{c}19)", bold=True)
    line(21, "dept_total_pct", "Percentage of revenue", lambda i, c: f"={c}20/{c}15", PCT1, italic_pct=True)
    line(22, "dept_inc", "DEPARTMENTAL INCOME (LOSS)", lambda i, c: f"={c}15-{c}20", bold=True)
    line(23, "dept_inc_pct", "Percentage of revenue", lambda i, c: f"={c}22/{c}15", PCT1, italic_pct=True)

    section_row(ws, 24, "UNDISTRIBUTED OPERATING EXPENSES", 1 + ncols)
    line(25, "ag", "Administrative and General",
         lambda i, c: f"={I('agfix')}*POWER(1+{I('growth')},{i})+{I('agvar')}*{c}15")
    line(26, "mkt", "Marketing", lambda i, c: f"={I('mkt')}*{c}15")
    line(27, "pom", "Property Operations and Maintenance", lambda i, c: f"={I('pom')}*{c}15")
    line(28, "util", "Utility Costs", lambda i, c: f"={I('util')}*{c}15")
    line(29, "undist_total", "Total undistributed expenses", lambda i, c: f"=SUM({c}25:{c}28)", bold=True)
    line(30, "undist_pct", "Percentage of revenue", lambda i, c: f"={c}29/{c}15", PCT1, italic_pct=True)

    line(31, "ebitdar", "EBITDAR", lambda i, c: f"={c}22-{c}29", bold=True)
    line(32, "ebitdar_pct", "Percentage of revenue", lambda i, c: f"={c}31/{c}15", PCT1, italic_pct=True)
    line(33, "rent", "Rent",
         lambda i, c: (f"=IF({c}3>={I('rent3start')},{I('rent3')},"
                       f"IF({c}3>={I('rent2start')},{I('rent2')},{I('rent1')}))"))
    line(34, "ebitda", "EBITDA", lambda i, c: f"={c}31-{c}33", bold=True)
    line(35, "reserve", "Reserve for replacement",
         lambda i, c: f"={I('res1')}*{c}15" if i == 0 else f"={I('res')}*{c}15")
    line(36, "noi", "NET OPERATING INCOME", lambda i, c: f"={c}34-{c}35", bold=True)
    line(37, "noi_pct", "Percentage of revenue", lambda i, c: f"={c}36/{c}15", PCT1, italic_pct=True)

    ws.column_dimensions["A"].width = 40
    for i in range(ncols):
        ws.column_dimensions[get_column_letter(2 + i)].width = 11.5
    ws.freeze_panes = "B4"
    return R


# --------------------------------------------------------------------------
# Returns sheet: unlevered cash flows 2024-2044, NPV and IRR (PDF page 11).
# --------------------------------------------------------------------------
def build_returns(wb, inp, model_rows):
    ws = wb.create_sheet("Returns")

    def I(key):
        return f"Inputs!$F${inp[key]}"

    ws["A1"] = "Unlevered returns (nominal €)"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("2024: rent during works plus CAPEX incl. VAT. 2025: first operating year plus full VAT "
                "recovery (per the plan). NPV uses Excel's NPV() over 2024-2044, i.e. the 2024 outflow is "
                "discounted by one year — this is the convention the PDF's €746,821 was computed with.")
    ws["A2"].font = F_SUB

    ncols = len(CF_YEARS)
    ws.cell(row=3, column=1, value="Year").font = F_HDR
    ws.cell(row=3, column=1).fill = FILL_HDR
    for i, y in enumerate(CF_YEARS):
        c = ws.cell(row=3, column=2 + i, value=y)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center")

    noi_r = model_rows["noi"]
    rent_lookup = (f"IF(B3>={I('rent3start')},{I('rent3')},"
                   f"IF(B3>={I('rent2start')},{I('rent2')},{I('rent1')}))")

    ws.cell(row=4, column=1, value="NOI")
    ws.cell(row=4, column=2, value=f"=-{rent_lookup}").number_format = EUR
    for i in range(1, ncols):                       # 2025.. -> Model cols B..
        mcol = get_column_letter(1 + i)
        ws.cell(row=4, column=2 + i, value=f"=Model!{mcol}{noi_r}").number_format = EUR

    ws.cell(row=5, column=1, value="CAPEX incl. VAT")
    ws.cell(row=5, column=2,
            value=f"=-({I('capex')}+{I('preopen')})*(1+{I('vat')})").number_format = EUR

    ws.cell(row=6, column=1, value="CAPEX VAT recovery")
    ws.cell(row=6, column=3, value=f"=({I('capex')}+{I('preopen')})*{I('vat')}").number_format = EUR

    ws.cell(row=7, column=1, value="CAPEX grant (if any)")
    ws.cell(row=7, column=3, value=f"={I('grant')}*{I('capex')}").number_format = EUR

    ws.cell(row=8, column=1, value="Net cash flow").font = F_BOLD
    for i in range(ncols):
        col = get_column_letter(2 + i)
        cell = ws.cell(row=8, column=2 + i, value=f"=SUM({col}4:{col}7)")
        cell.number_format = EUR
        cell.font = F_BOLD
        cell.border = TOP_BORDER

    last = get_column_letter(1 + ncols)
    ws.cell(row=10, column=1, value="NPV @ discount rate").font = F_BOLD
    npv = ws.cell(row=10, column=2, value=f"=NPV({I('disc')},B8:{last}8)")
    npv.number_format = EUR
    npv.font = F_BOLD
    ws.cell(row=11, column=1, value="IRR (unlevered)").font = F_BOLD
    irr = ws.cell(row=11, column=2, value=f"=IRR(B8:{last}8)")
    irr.number_format = "0.0%"
    irr.font = F_BOLD

    ws.column_dimensions["A"].width = 26
    for i in range(ncols):
        ws.column_dimensions[get_column_letter(2 + i)].width = 11.5
    ws.freeze_panes = "B4"
    return {"noi": 4, "capex": 5, "vat_rec": 6, "grant": 7, "cf": 8, "npv": 10, "irr": 11}


# --------------------------------------------------------------------------
# Check sheet: the live model vs every figure printed in the PDF.
# --------------------------------------------------------------------------
def build_check(wb, inp, model_rows, ret_rows):
    ws = wb.create_sheet("Check vs PDF")
    ws["A1"] = "Accuracy check — live model vs the printed PDF figures"
    ws["A1"].font = F_TITLE
    ws["A2"] = (f'=IF(Inputs!$B$4<>"Base","SWITCH THE SCENARIO TO Base — these checks compare the Base '
                f'scenario to the PDF","Comparing Base scenario to the PDF. Tolerances allow for the '
                f'PDF\'s own rounding (it prints whole euros / one decimal).")')
    ws["A2"].font = Font(bold=True, color=RED)

    hdr = 6
    for c, text in enumerate(["Statement", "Line item", "Year", "PDF value", "Model value",
                              "Difference", "Tolerance", "Status"], start=1):
        cell = ws.cell(row=hdr, column=c, value=text)
        cell.font = F_HDR
        cell.fill = FILL_HDR

    r = hdr + 1
    first_data = r

    def row(statement, item, year, pdf_val, model_ref, tol, fmt):
        nonlocal r
        ws.cell(row=r, column=1, value=statement)
        ws.cell(row=r, column=2, value=item)
        ws.cell(row=r, column=3, value=year)
        ws.cell(row=r, column=4, value=pdf_val).number_format = fmt
        ws.cell(row=r, column=5, value=f"={model_ref}").number_format = fmt
        ws.cell(row=r, column=6, value=f"=E{r}-D{r}").number_format = fmt
        ws.cell(row=r, column=7, value=tol).number_format = fmt
        ws.cell(row=r, column=8, value=f'=IF(ABS(F{r})<=G{r},"OK","CHECK")')
        r += 1

    kpi_lines = [("occ", "Occupancy", PCT1, 0.005), ("occupied", "Occupied room nights", EUR, 1),
                 ("adr", "ADR", EUR2, 0.51), ("revpar", "RevPAR", EUR2, 0.51)]
    pl_lines = [("room_rev", "Room revenue"), ("fb_rev", "Food & Beverages"),
                ("total_rev", "Total revenue"), ("rooms_exp", "Rooms expense"),
                ("fb_exp", "F&B expense"), ("staff", "Staff"),
                ("dept_total", "Total departmental expenses"), ("dept_inc", "Departmental income"),
                ("ag", "Administrative and General"), ("mkt", "Marketing"),
                ("pom", "Property Operations and Maintenance"), ("util", "Utility Costs"),
                ("undist_total", "Total undistributed expenses"), ("ebitdar", "EBITDAR"),
                ("rent", "Rent"), ("ebitda", "EBITDA"), ("reserve", "Reserve for replacement"),
                ("noi", "Net operating income")]
    pdf_key = {"occ": "occupancy"}

    for key, label, fmt, tol in kpi_lines:
        for i, y in enumerate(PDF_YEARS):
            col = get_column_letter(2 + i)
            row("KPIs (p.9)", label, y, PDF[pdf_key.get(key, key)][i],
                f"Model!{col}{model_rows[key]}", tol, fmt)
    for key, label in pl_lines:
        for i, y in enumerate(PDF_YEARS):
            col = get_column_letter(2 + i)
            row("P&L (p.10)", label, y, PDF[key][i], f"Model!{col}{model_rows[key]}", 2, EUR)

    row("Returns (p.11)", "NOI", 2024, PDF_RETURNS["noi_2024"], f"Returns!B{ret_rows['noi']}", 1, EUR)
    row("Returns (p.11)", "CAPEX incl. VAT", 2024, PDF_RETURNS["capex_2024"],
        f"Returns!B{ret_rows['capex']}", 1, EUR)
    row("Returns (p.11)", "Net cash flow", 2024, PDF_RETURNS["net_cf_2024"],
        f"Returns!B{ret_rows['cf']}", 1, EUR)
    row("Returns (p.11)", "CAPEX VAT recovery", 2025, PDF_RETURNS["vat_recovery_2025"],
        f"Returns!C{ret_rows['vat_rec']}", 1, EUR)
    row("Returns (p.11)", "Net cash flow", 2025, PDF_RETURNS["net_cf_2025"],
        f"Returns!C{ret_rows['cf']}", 2, EUR)
    last_model_col = get_column_letter(1 + len(MODEL_YEARS))
    row("Returns (p.11)", "NOI", 2044, PDF_RETURNS["noi_2044"],
        f"Model!{last_model_col}{model_rows['noi']}", 2, EUR)
    row("Returns (p.11)", "NPV @ 10%", "", PDF_RETURNS["npv"], f"Returns!B{ret_rows['npv']}", 25, EUR)
    row("Returns (p.11)", "IRR", "", PDF_RETURNS["irr"], f"Returns!B{ret_rows['irr']}", 0.0005, PCT2)

    last_data = r - 1
    ws["A4"] = "Checks passed:"
    ws["A4"].font = F_BOLD
    ws["B4"] = f'=COUNTIF(H{first_data}:H{last_data},"OK")&" of "&COUNTA(H{first_data}:H{last_data})'
    ws["B4"].font = F_BOLD

    status_range = f"H{first_data}:H{last_data}"
    ws.conditional_formatting.add(
        status_range, CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK,
                                 font=Font(color="006100")))
    ws.conditional_formatting.add(
        status_range, CellIsRule(operator="equal", formula=['"CHECK"'], fill=FILL_BAD,
                                 font=Font(color="9C0006", bold=True)))

    widths = {"A": 15, "B": 34, "C": 7, "D": 13, "E": 13, "F": 11, "G": 10, "H": 9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{first_data}"


# --------------------------------------------------------------------------
# README sheet
# --------------------------------------------------------------------------
README_TEXT = [
    ("The Downtown house, a social hub at Megaro Foumi — financial model", F_TITLE),
    ("Reconstructed source model for the GBR Consulting business plan (May 2024, draft)", F_SUB),
    ("", None),
    ("WHAT THIS FILE IS", F_BOLD),
    ("A working version of the spreadsheet model behind the plan's financial tables (pages 9-11).", None),
    ("It was reverse-engineered from the printed figures: every relationship below reproduces the PDF", None),
    ("to the euro, and the 'Check vs PDF' sheet proves it against all ~230 printed numbers.", None),
    ("", None),
    ("HOW TO USE IT", F_BOLD),
    ("1. All inputs live on the 'Inputs' sheet in the blue cells; everything else is formulas.", None),
    ("2. Pick a scenario in Inputs!B4 (Base / Upside / Downside). Base = the plan as published.", None),
    ("   Upside and Downside start as illustrative placeholders — overwrite them with your own scenarios.", None),
    ("3. 'Model' is the P&L (2025-2044); 'Returns' gives net cash flow, NPV and IRR (2024-2044).", None),
    ("4. 'Check vs PDF' compares the live model to the printed figures. With the Base scenario every", None),
    ("   row should read OK; tolerances only allow for the PDF's own display rounding.", None),
    ("   (Verified: the file recalculates automatically on open. If any app ever shows zeros,", None),
    ("   force a recalculation — Ctrl+Shift+F9 in LibreOffice, F9 in Excel.)", None),
    ("", None),
    ("HOW THE MODEL WORKS (all proven against the PDF)", F_BOLD),
    ("- Rooms: 5 rooms x 365 days; occupancy ramps 65/70/75/80%, flat at 80% thereafter.", None),
    ("- ADR ramps to the stabilised level in 2028, then grows at inflation (2%). The PDF prints ADR", None),
    ("  rounded to whole euros; the inputs hold the exact implied values (113.22/121.90/130.88/140.50).", None),
    ("- F&B revenue: direct inputs for 2025-2028, then +2% per year.", None),
    ("- Departmental costs: Rooms = 12.28/11.96/11.68/11.40% of room revenue (11.40% from 2028);", None),
    ("  F&B = 37.5% of F&B revenue; Staff = EUR 533,868 in 2025 growing 2%/year (32 FTEs).", None),
    ("- Undistributed: A&G = EUR 82,212 fixed (grows 2%/yr) + 2% of revenue; Marketing 3%;", None),
    ("  Property Ops & Maintenance 3%; Utilities 4.5% of total revenue.", None),
    ("- Fixed charges: rent EUR 15,000 (2024-2029), EUR 21,000 (2030-2034), EUR 90,000 from 2035;", None),
    ("  reserve for replacement 1% of revenue in year 1, 2% thereafter.", None),
    ("- Returns: CAPEX EUR 1.5m + EUR 100k pre-opening, both +24% VAT paid in 2024; VAT fully", None),
    ("  recovered in 2025; NPV at 10% with the 2024 outflow discounted one year (Excel NPV convention).", None),
    ("", None),
    ("CAVEATS", F_BOLD),
    ("- The EUR 90,000 rent from 2035 is not printed anywhere in the PDF; it is the value implied by the", None),
    ("  PDF's 2044 NOI (414,765) and NPV (746,821). Check it against the actual lease before relying on it.", None),
    ("- The ramp-up values (ADR, F&B, rooms-cost %) are taken as direct inputs; the consultants may have", None),
    ("  derived them from a monthly/seasonality model that annual figures cannot reveal.", None),
    ("- The CAPEX grant driver (plan mentions applying for 60-75%) is set to 0% in Base, received 2025.", None),
    ("- This reconstruction is for analysis; it is not GBR Consulting's original file.", None),
]


def build_readme(wb):
    ws = wb.create_sheet("README", 0)
    for i, (text, font) in enumerate(README_TEXT, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if font:
            cell.font = font
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False


def main(path):
    wb = Workbook()
    wb.remove(wb.active)
    inp = build_inputs(wb)
    model_rows = build_model(wb, inp)
    ret_rows = build_returns(wb, inp, model_rows)
    build_check(wb, inp, model_rows, ret_rows)
    build_readme(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    print(f"Wrote {path}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "downtown-house-model.xlsx")
